"""Parse uploaded documents (PDF, DOCX, Markdown, plain text) into text."""

import logging
from pathlib import Path

from app.services.file_storage import resolve_file_path
from app.services.pdf_service import parse_pdf

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".json",
    ".md",
    ".txt",
    ".fasta",
    ".fa",
    ".faa",
}


async def parse_document(
    file_path: str,
    *,
    allow_unmanaged_absolute: bool = True,
) -> str:
    """Parse a document file and return its text content.

    Dispatches on the file extension:
    - .pdf -> marker-based PDF parsing
    - .docx -> python-docx paragraphs and table cells
    - .json / .md / .txt / .fasta / .fa / .faa -> plain text with utf-8/gbk decoding fallback
    """
    resolved = str(
        resolve_file_path(
            file_path,
            allow_legacy_absolute=allow_unmanaged_absolute,
        )
    )
    ext = Path(resolved).suffix.lower()

    if ext == ".pdf":
        return await parse_pdf(resolved)
    if ext == ".docx":
        return _parse_docx(resolved)
    if ext in {".json", ".md", ".txt", ".fasta", ".fa", ".faa"}:
        return _parse_text(resolved)
    raise ValueError(f"Unsupported file type: {ext or Path(resolved).name}")


def _parse_docx(file_path: str) -> str:
    """Extract paragraph and table cell text from a DOCX file."""
    from docx import Document

    doc = Document(file_path)
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)


def _parse_text(file_path: str) -> str:
    """Read a plain-text file, decoding as utf-8 with gbk fallback."""
    raw = Path(file_path).read_bytes()
    for encoding in ("utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")
