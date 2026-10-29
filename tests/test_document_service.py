import io
import zipfile

import pytest

from app.services import document_service
from app.services.document_service import parse_document
from app.services.zip_service import extract_documents_from_zip

pytestmark = pytest.mark.asyncio


async def test_parse_markdown_utf8(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Title\n\nSome **markdown** content.", encoding="utf-8")

    text = await parse_document(str(path))

    assert "# Title" in text
    assert "markdown" in text


async def test_parse_txt_gbk_fallback(tmp_path):
    path = tmp_path / "gbk-notes.txt"
    path.write_bytes("中文实验记录".encode("gbk"))

    text = await parse_document(str(path))

    assert "中文实验记录" in text


async def test_parse_fasta_as_plain_text(tmp_path):
    path = tmp_path / "proteins.fasta"
    path.write_text(">protein_001\nMKKIAVIGTASLAVATASAQ\n", encoding="utf-8")

    text = await parse_document(str(path))

    assert text == ">protein_001\nMKKIAVIGTASLAVATASAQ\n"


async def test_parse_json_as_plain_text(tmp_path):
    path = tmp_path / "allele-panel.json"
    path.write_text('{"schema_version": 1}', encoding="utf-8")

    text = await parse_document(str(path))

    assert text == '{"schema_version": 1}'


async def test_parse_txt_replaces_undecodable_bytes(tmp_path):
    path = tmp_path / "broken.txt"
    path.write_bytes(b"valid \xff\xff invalid")

    text = await parse_document(str(path))

    assert "valid" in text
    assert "invalid" in text


async def test_parse_docx_extracts_paragraphs_and_tables(tmp_path):
    from docx import Document

    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    doc.save(str(path))

    text = await parse_document(str(path))

    assert "First paragraph" in text
    assert "Second paragraph" in text
    assert "Cell A" in text
    assert "Cell B" in text


async def test_parse_pdf_delegates_to_parse_pdf(monkeypatch, tmp_path):
    path = tmp_path / "paper.pdf"
    path.write_bytes(b"%PDF-1.4\n%%EOF")
    called = []

    async def fake_parse_pdf(file_path: str) -> str:
        called.append(file_path)
        return "pdf text"

    monkeypatch.setattr(document_service, "parse_pdf", fake_parse_pdf)

    text = await parse_document(str(path))

    assert text == "pdf text"
    assert called == [str(path)]


async def test_parse_document_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        await parse_document(str(path))


def _make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buffer.getvalue()


def test_zip_extraction_accepts_new_document_types():
    archive = _make_zip(
        {
            "paper.pdf": b"%PDF-1.4\n%%EOF",
            "report.docx": b"docx-bytes",
            "notes.md": b"# notes",
            "readme.txt": b"plain text",
            "data.csv": b"a,b\n1,2\n",
        }
    )

    documents, errors = extract_documents_from_zip(archive)

    assert sorted(d.filename for d in documents) == [
        "notes.md",
        "paper.pdf",
        "readme.txt",
        "report.docx",
    ]
    assert errors == []


def test_zip_extraction_still_skips_unsupported_and_hidden_files():
    archive = _make_zip(
        {
            "paper.pdf": b"%PDF-1.4\n%%EOF",
            "image.png": b"png-bytes",
            ".hidden.pdf": b"%PDF-1.4\n%%EOF",
            "__MACOSX/._paper.pdf": b"junk",
        }
    )

    documents, errors = extract_documents_from_zip(archive)

    assert [d.filename for d in documents] == ["paper.pdf"]
    assert errors == []
